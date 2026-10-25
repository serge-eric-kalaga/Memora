"""Read-only, in-memory exports. Source text is data, never HTML or executable input."""

from datetime import datetime, UTC
from io import BytesIO
import math
from pathlib import Path
import re
import threading
import unicodedata
from urllib.parse import quote
from xml.sax.saxutils import escape

FORMATS = {
    "transcript": {"txt", "pdf", "docx", "srt", "vtt"},
    "report": {"txt", "pdf", "docx", "md"},
}
MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "srt": "application/x-subrip; charset=utf-8",
    "vtt": "text/vtt; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_font_lock = threading.Lock()


def timestamp(seconds, milliseconds=False, separator="."):
    value = float(seconds or 0)
    if not math.isfinite(value):
        value = 0
    total = max(0, round(value * 1000))
    whole, ms = divmod(total, 1000)
    hours, remainder = divmod(whole, 3600)
    minutes, secs = divmod(remainder, 60)
    base = f"{hours:02}:{minutes:02}:{secs:02}"
    return base + f"{separator}{ms:03}" if milliseconds else base


def created_at(value):
    return datetime.fromtimestamp(value, UTC).strftime("%d/%m/%Y %H:%M UTC")


def disposition(title, content, revision, extension):
    # Both names are bounded and stripped of header/path control characters.
    title = re.sub(r"[\x00-\x1f\x7f/\\]+", "-", str(title)).strip(" .-")[:90]
    title = title or "Memora"
    kind = "transcription" if content == "transcript" else "rapport"
    filename = f"{title}-{kind}-v{revision}.{extension}"
    ascii_name = (
        unicodedata.normalize("NFKD", filename).encode("ascii", "ignore").decode()
    )
    ascii_name = re.sub(r"[^a-zA-Z0-9._ -]", "-", ascii_name)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename, safe='')}"


def normalized_report(report):
    """Older stored reports may use plain strings instead of structured items."""
    if not isinstance(report, dict):
        return {}
    result = {"limitations": str(report.get("limitations") or "")}
    for key in ("summary", "topics", "decisions", "actions", "open_questions"):
        values = report.get(key) or []
        if isinstance(values, str):
            values = [values]
        if not isinstance(values, list):
            values = []
        result[key] = []
        for value in values:
            item = {"text": value} if isinstance(value, str) else value
            if isinstance(item, dict):
                result[key].append(
                    {
                        **item,
                        "sources": [
                            source
                            for source in item.get("sources", [])
                            if isinstance(source, str)
                        ]
                        if isinstance(item.get("sources", []), list)
                        else [],
                    }
                )
    return result


def blocks(snapshot):
    """Simple document model shared by TXT, Markdown, PDF and DOCX."""
    meeting, content = snapshot["meeting"], snapshot["content"]
    result = [("title", meeting["title"])]
    version = snapshot["version"]
    historical = snapshot["historical"]
    kind = "Transcription" if content == "transcript" else "Rapport"
    result.append(
        (
            "meta",
            f"{kind} · Version {snapshot['revision']} · Transcription source v{version}",
        )
    )
    result.append(
        (
            "meta",
            f"Créé le {created_at(snapshot['created'])} · {'Version historique' if historical else 'Version courante'}",
        )
    )
    if content == "transcript":
        for seg in snapshot["segments"]:
            label = seg.get("speaker") or "Intervenant"
            uncertainty = " · Passage incertain" if seg.get("uncertain") else ""
            result.append(
                (
                    "heading",
                    f"{timestamp(seg.get('start'))} - {timestamp(seg.get('end'))} · {label}{uncertainty}",
                )
            )
            result.append(("text", seg.get("text", "")))
        return result
    if snapshot["stale"]:
        result.append(
            (
                "notice",
                "À actualiser : ce rapport ne correspond plus à la transcription courante. Les révisions humaines ci-dessous concernent cette version.",
            )
        )
    result.append(
        (
            "notice",
            "Compte rendu généré par IA. Vérifiez les éléments importants avec les sources disponibles.",
        )
    )
    reviews = {(r["category"], r["item_index"]): r for r in snapshot["reviews"]}
    for key, label in [
        ("summary", "Résumé"),
        ("topics", "Sujets abordés"),
        ("decisions", "Décisions"),
        ("actions", "Actions à suivre"),
        ("open_questions", "Questions ouvertes"),
    ]:
        result.append(("section", label))
        items = snapshot["report"].get(key, [])
        if not items:
            result.append(("text", "Aucun élément établi dans la transcription."))
        for index, item in enumerate(items):
            result.append(("text", item.get("text", "")))
            if key == "actions":
                result.append(
                    (
                        "meta",
                        f"Responsable : {item.get('owner') or 'Non précisé'} · Échéance : {item.get('deadline') or 'Non précisée'}",
                    )
                )
            if key in ("decisions", "actions", "open_questions"):
                refs = [
                    snapshot["source_map"][sid]
                    for sid in item.get("sources", [])
                    if sid in snapshot["source_map"]
                ]
                if refs:
                    result.append(("meta", "Sources : " + "; ".join(refs)))
            if key in ("decisions", "actions"):
                review = reviews.get((key, index))
                status = {
                    "proposed": "Proposé",
                    "confirmed": "Validé",
                    "needs_review": "À vérifier",
                    "cancelled": "Annulé",
                }.get((review or {}).get("status"), "Non relu")
                result.append(("meta", "Révision humaine : " + status))
                if review:
                    result.append(
                        (
                            "meta",
                            f"Révision humaine n°{review['revision']} · {created_at(review['created'])}",
                        )
                    )
                    for field, name in [
                        ("text", "Texte"),
                        ("owner", "Responsable"),
                        ("deadline", "Échéance"),
                    ]:
                        if field in review.get("corrections", {}):
                            result.append(
                                (
                                    "text",
                                    f"Correction humaine — {name} : {review['corrections'][field]}",
                                )
                            )
    if snapshot["report"].get("limitations"):
        result.extend(
            [
                ("section", "Points à vérifier"),
                ("text", snapshot["report"]["limitations"]),
            ]
        )
    return result


def subtitle(snapshot, extension):
    lines = ["WEBVTT", ""] if extension == "vtt" else []
    for index, seg in enumerate(snapshot["segments"], start=1):
        start = max(0, float(seg.get("start", 0)))
        end = max(start + 0.001, float(seg.get("end", start)))
        separator = "," if extension == "srt" else "."
        lines.extend(
            [
                str(index),
                f"{timestamp(start, True, separator)} --> {timestamp(end, True, separator)}",
            ]
        )
        speaker = (
            str(seg.get("speaker") or "Intervenant")
            .replace("\n", " ")
            .replace("\r", " ")
        )
        text = str(seg.get("text", "")).replace("\r", " ").replace("\n", " ")
        # Escape WebVTT markup so user text cannot change cue rendering.
        cue = f"{speaker} : {text}"
        if seg.get("uncertain"):
            cue += " [Passage incertain]"
        cue = escape(cue)
        lines.extend([cue, ""])
    return ("\n".join(lines) + "\n").encode("utf-8", errors="backslashreplace")


def document_entries(entries, glyphs=None):
    """Represent unsupported characters explicitly rather than silently dropping them."""
    changed = False
    cleaned = []
    for kind, text in entries:
        value = []
        for char in str(text):
            code = ord(char)
            xml_valid = (
                code in (9, 10, 13)
                or 32 <= code <= 0xD7FF
                or 0xE000 <= code <= 0xFFFD
                or 0x10000 <= code <= 0x10FFFF
            )
            if not xml_valid or (
                glyphs is not None and char not in "\n\r\t" and code not in glyphs
            ):
                value.append(f"[U+{code:04X}]")
                changed = True
            else:
                value.append(char)
        cleaned.append((kind, "".join(value)))
    if changed:
        cleaned.insert(
            1,
            (
                "notice",
                "Certains caractères non représentables dans ce format sont indiqués par leur code [U+XXXX]. Le format TXT conserve le texte Unicode original.",
            ),
        )
    return cleaned


def pdf_document(entries):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    with _font_lock:
        if "MemoraVera" not in pdfmetrics.getRegisteredFontNames():
            fonts = Path(reportlab.__file__).parent / "fonts"
            pdfmetrics.registerFont(TTFont("MemoraVera", str(fonts / "Vera.ttf")))
            pdfmetrics.registerFont(TTFont("MemoraVeraBold", str(fonts / "VeraBd.ttf")))
    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        title="Memora",
        author="Memora",
        leftMargin=48,
        rightMargin=48,
        topMargin=45,
        bottomMargin=48,
    )
    specs = {
        "title": (20, 26, "MemoraVeraBold", "#232336"),
        "section": (14, 20, "MemoraVeraBold", "#6550df"),
        "heading": (10, 15, "MemoraVeraBold", "#414153"),
        "text": (10, 16, "MemoraVera", "#232336"),
        "meta": (8, 12, "MemoraVera", "#696a80"),
        "notice": (9, 14, "MemoraVera", "#696a80"),
    }
    styles = {
        kind: ParagraphStyle(
            kind,
            fontName=font,
            fontSize=size,
            leading=leading,
            textColor=colors.HexColor(color),
            alignment=TA_LEFT,
            splitLongWords=True,
            spaceAfter=7,
            keepWithNext=kind in ("title", "section", "heading"),
        )
        for kind, (size, leading, font, color) in specs.items()
    }
    story = []
    for kind, text in document_entries(
        entries, pdfmetrics.getFont("MemoraVera").face.charToGlyph
    ):
        # Paragraph supports HTML; escape every source string before rendering.
        for part in str(text).split("\n"):
            story.append(Paragraph(escape(part) or "&#160;", styles[kind]))
        if kind == "section":
            story.append(Spacer(1, 3))

    def footer(canvas, document):
        canvas.setFont("MemoraVera", 8)
        canvas.setFillColor(colors.HexColor("#696a80"))
        canvas.drawString(48, 28, "Memora")
        canvas.drawRightString(document.pagesize[0] - 48, 28, f"Page {document.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def docx_document(entries):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.core_properties.author = "Memora"
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Title"].font.color.rgb = RGBColor.from_string("000000")
    for style_name in ("Title", "Normal"):
        style = document.styles[style_name]
        for borders in list(style.element.iter(qn("w:pBdr"))):
            borders.getparent().remove(borders)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.color.rgb = RGBColor.from_string("6550DF")
    for kind, text in document_entries(entries):
        if kind == "title":
            title = document.add_heading(str(text), 0)
            for borders in list(title._p.iter(qn("w:pBdr"))):
                borders.getparent().remove(borders)
        elif kind in ("section", "heading"):
            document.add_heading(str(text), 1 if kind == "section" else 2)
        else:
            paragraph = document.add_paragraph(str(text))
            if kind in ("meta", "notice"):
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string("696A80")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Memora · Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def render(snapshot, extension):
    if extension in ("srt", "vtt"):
        return subtitle(snapshot, extension)
    entries = blocks(snapshot)
    if extension == "pdf":
        return pdf_document(entries)
    if extension == "docx":
        return docx_document(entries)
    if extension == "md":
        return (
            "\n\n".join(
                (
                    "# "
                    if kind == "title"
                    else "## "
                    if kind == "section"
                    else "### "
                    if kind == "heading"
                    else ""
                )
                # Source content stays literal in Markdown viewers too: no raw
                # HTML, image loads, links or nested formatting from the audio.
                + re.sub(r"([\\`*_{}\[\]()#+\-.!>|~])", r"\\\1", escape(str(text)))
                for kind, text in entries
            )
            + "\n"
        ).encode("utf-8", errors="backslashreplace")
    return ("\n\n".join(str(text) for _, text in entries) + "\n").encode(
        "utf-8", errors="backslashreplace"
    )
